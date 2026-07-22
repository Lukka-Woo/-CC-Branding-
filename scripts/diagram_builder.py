"""
Diagram generation utilities for PPTX slides.

Orthogonal arrows with rounded elbows are drawn as a single custom-geometry
path (custGeom) using cubic Bézier curves to approximate quarter-circle arcs.
This avoids the hard corners produced by chaining straight connectors.

All coordinate arguments accept EMU integers (e.g. Mm(21)).
Color arguments accept 6-char hex strings without '#' (matching BT constants).
"""

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

from scripts.pptx_builder import _apply_radius, _rgb, _txb, _rect, _card
import scripts.brand_tokens as BT

# Cubic Bézier approximation constant for a quarter-circle arc
_K = 0.5523

# XML namespace helpers
_PNS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
_ANS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

def _p(tag): return f'{{{_PNS}}}{tag}'
def _a(tag): return f'{{{_ANS}}}{tag}'


# ─────────────────────────────────────────────────────────────────────────────
# Internal: custom-geometry path builder
# ─────────────────────────────────────────────────────────────────────────────

def _next_shape_id(slide):
    ids = [int(el.get('id', 0))
           for el in slide.shapes._spTree.iter(qn('p:cNvPr'))
           if el.get('id')]
    return max(ids, default=0) + 1


def _make_path_shape(slide, rel_cmds, box_x, box_y, box_w, box_h,
                     color_hex, width_pt, style, head):
    """
    Insert a <p:sp> with custGeom into the slide's shape tree.
    rel_cmds: list of ('moveTo',x,y) | ('lnTo',x,y) | ('cubicBezTo',(x1,y1),(x2,y2),(x3,y3))
              with coordinates relative to (box_x, box_y).
    """
    shape_id   = _next_shape_id(slide)
    line_w_emu = str(int(Pt(width_pt)))
    color_rgb  = color_hex.upper().lstrip('#')

    sp = etree.Element(_p('sp'))

    # Non-visual props
    nvSpPr = etree.SubElement(sp, _p('nvSpPr'))
    cNvPr  = etree.SubElement(nvSpPr, _p('cNvPr'))
    cNvPr.set('id', str(shape_id))
    cNvPr.set('name', f'diag{shape_id}')
    cNvSpPr = etree.SubElement(nvSpPr, _p('cNvSpPr'))
    etree.SubElement(cNvSpPr, _a('spLocks')).set('noGrp', '1')
    etree.SubElement(nvSpPr, _p('nvPr'))

    # Shape properties
    spPr = etree.SubElement(sp, _p('spPr'))
    xfrm = etree.SubElement(spPr, _a('xfrm'))
    off  = etree.SubElement(xfrm, _a('off'))
    off.set('x', str(int(box_x))); off.set('y', str(int(box_y)))
    ext  = etree.SubElement(xfrm, _a('ext'))
    ext.set('cx', str(int(box_w))); ext.set('cy', str(int(box_h)))

    # Custom geometry
    cg  = etree.SubElement(spPr, _a('custGeom'))
    etree.SubElement(cg, _a('avLst'))
    etree.SubElement(cg, _a('gdLst'))
    etree.SubElement(cg, _a('ahLst'))
    etree.SubElement(cg, _a('cxnLst'))
    rect = etree.SubElement(cg, _a('rect'))
    rect.set('l', '0'); rect.set('t', '0'); rect.set('r', '0'); rect.set('b', '0')
    pathLst = etree.SubElement(cg, _a('pathLst'))
    path_el = etree.SubElement(pathLst, _a('path'))
    path_el.set('w', str(int(box_w))); path_el.set('h', str(int(box_h)))
    path_el.set('fill', 'none')

    for cmd in rel_cmds:
        if cmd[0] == 'moveTo':
            mt = etree.SubElement(path_el, _a('moveTo'))
            pt = etree.SubElement(mt, _a('pt'))
            pt.set('x', str(cmd[1])); pt.set('y', str(cmd[2]))
        elif cmd[0] == 'lnTo':
            lt = etree.SubElement(path_el, _a('lnTo'))
            pt = etree.SubElement(lt, _a('pt'))
            pt.set('x', str(cmd[1])); pt.set('y', str(cmd[2]))
        elif cmd[0] == 'cubicBezTo':
            bz = etree.SubElement(path_el, _a('cubicBezTo'))
            for px, py in cmd[1:]:
                pt = etree.SubElement(bz, _a('pt'))
                pt.set('x', str(px)); pt.set('y', str(py))

    # No fill on shape
    etree.SubElement(spPr, _a('noFill'))

    # Line style
    ln = etree.SubElement(spPr, _a('ln'))
    ln.set('w', line_w_emu)
    sf = etree.SubElement(ln, _a('solidFill'))
    sc = etree.SubElement(sf, _a('srgbClr'))
    sc.set('val', color_rgb)
    if style != 'solid':
        pd = etree.SubElement(ln, _a('prstDash'))
        pd.set('val', style)
    if head in ('triangle', 'arrow'):
        te = etree.SubElement(ln, _a('tailEnd'))
        te.set('type', head); te.set('w', 'med'); te.set('len', 'med')

    # Minimal style / txBody (required by spec)
    st = etree.SubElement(sp, _p('style'))
    for tag, idx, clr in [('lnRef','2','accent1'),('fillRef','1','accent1'),
                           ('effectRef','0','accent1'),('fontRef','1','lt1')]:
        ref = etree.SubElement(st, _a(tag))
        ref.set('idx', idx)
        etree.SubElement(ref, _a('schemeClr')).set('val', clr)
    txBody = etree.SubElement(sp, _p('txBody'))
    etree.SubElement(txBody, _a('bodyPr'))
    etree.SubElement(txBody, _a('lstStyle'))
    etree.SubElement(txBody, _a('p'))

    slide.shapes._spTree.append(sp)
    return sp


def _build_rounded_cmds(waypoints, corner_mm):
    """
    Convert a list of orthogonal waypoints into path commands with rounded
    corners (quarter-circle Bézier arcs).

    Returns (abs_cmds, min_x, min_y, max_x, max_y) where abs_cmds uses
    absolute EMU coordinates.
    """
    r = Mm(corner_mm)
    n = len(waypoints)
    abs_cmds = []

    abs_cmds.append(('moveTo', int(waypoints[0][0]), int(waypoints[0][1])))

    for i in range(1, n):
        pt   = waypoints[i]
        prev = waypoints[i - 1]
        dx = pt[0] - prev[0]; dy = pt[1] - prev[1]
        seg_len = abs(dx) + abs(dy)
        sx = (1 if dx > 0 else -1) if dx != 0 else 0
        sy = (1 if dy > 0 else -1) if dy != 0 else 0

        has_next = (i < n - 1)

        if has_next:
            nxt = waypoints[i + 1]
            ndx = nxt[0] - pt[0]; ndy = nxt[1] - pt[1]
            next_len = abs(ndx) + abs(ndy)
            nsx = (1 if ndx > 0 else -1) if ndx != 0 else 0
            nsy = (1 if ndy > 0 else -1) if ndy != 0 else 0

            actual_r = min(r, seg_len // 2, next_len // 2)

            # Line to (corner − r along incoming direction)
            lx = int(pt[0] - sx * actual_r)
            ly = int(pt[1] - sy * actual_r)
            abs_cmds.append(('lnTo', lx, ly))

            # Bézier arc at corner
            p3x = int(pt[0] + nsx * actual_r)
            p3y = int(pt[1] + nsy * actual_r)
            p1x = int(lx + _K * actual_r * sx)
            p1y = int(ly + _K * actual_r * sy)
            p2x = int(p3x - _K * actual_r * nsx)
            p2y = int(p3y - _K * actual_r * nsy)
            abs_cmds.append(('cubicBezTo', (p1x, p1y), (p2x, p2y), (p3x, p3y)))
        else:
            abs_cmds.append(('lnTo', int(pt[0]), int(pt[1])))

    # Bounding box
    all_x, all_y = [], []
    for cmd in abs_cmds:
        if cmd[0] in ('moveTo', 'lnTo'):
            all_x.append(cmd[1]); all_y.append(cmd[2])
        elif cmd[0] == 'cubicBezTo':
            for px, py in cmd[1:]:
                all_x.append(px); all_y.append(py)

    PAD = max(int(r * 0.15), int(Mm(0.8)))
    return abs_cmds, min(all_x) - PAD, min(all_y) - PAD, max(all_x) + PAD, max(all_y) + PAD


def _orth_arrow_rounded(slide, waypoints, color, width_pt, style, head, corner_mm):
    abs_cmds, min_x, min_y, max_x, max_y = _build_rounded_cmds(waypoints, corner_mm)
    box_w = max_x - min_x
    box_h = max_y - min_y

    def rel(x, y): return int(x - min_x), int(y - min_y)

    rel_cmds = []
    for cmd in abs_cmds:
        if cmd[0] in ('moveTo', 'lnTo'):
            rel_cmds.append((cmd[0], *rel(cmd[1], cmd[2])))
        elif cmd[0] == 'cubicBezTo':
            rel_cmds.append(('cubicBezTo', rel(*cmd[1]), rel(*cmd[2]), rel(*cmd[3])))

    _make_path_shape(slide, rel_cmds, min_x, min_y, box_w, box_h,
                     color, width_pt, style, head)


# ─────────────────────────────────────────────────────────────────────────────
# Internal: simple multi-connector (fallback for 2-point or corner_mm=0)
# ─────────────────────────────────────────────────────────────────────────────

def _conn(slide, x1, y1, x2, y2):
    from pptx.enum.shapes import MSO_CONNECTOR_TYPE
    return slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT, int(x1), int(y1), int(x2), int(y2))


def _style_line(conn, color_hex, width_pt, style='solid', head=None):
    conn.line.color.rgb = _rgb(color_hex)
    conn.line.width = Pt(width_pt)
    cxn_sp = conn._element
    spPr = cxn_sp.find(qn('p:spPr'))
    if spPr is None: return
    ln = spPr.find(qn('a:ln'))
    if ln is None: return
    if style != 'solid':
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', style)
    if head in ('triangle', 'arrow'):
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', head); te.set('w', 'med'); te.set('len', 'med')


def _orth_arrow_simple(slide, waypoints, color, width_pt, style, head):
    connectors = []
    n = len(waypoints)
    for i in range(n - 1):
        x1, y1 = waypoints[i]; x2, y2 = waypoints[i + 1]
        is_last = (i == n - 2)
        c = _conn(slide, x1, y1, x2, y2)
        _style_line(c, color, width_pt, style=style,
                    head=head if is_last else None)
        connectors.append(c)
    return connectors


# ─────────────────────────────────────────────────────────────────────────────
# Public: orthogonal arrows
# ─────────────────────────────────────────────────────────────────────────────

def orth_arrow(slide, waypoints, color, width_pt=1.2,
               style='solid', head='triangle', corner_mm=3.0):
    """
    Draw an orthogonal (axis-aligned) arrow through waypoints.

    Consecutive points must share x OR y (no diagonals).
    corner_mm > 0 produces rounded elbows via Bézier arcs.
    style: 'solid' | 'dash' | 'sysDash' | 'lgDash' | 'dashDot'
    head:  'triangle' | 'arrow' | 'none'
    """
    n = len(waypoints)
    if n < 2:
        return
    if n == 2 or corner_mm <= 0:
        _orth_arrow_simple(slide, waypoints, color, width_pt, style,
                           head if head != 'none' else None)
    else:
        _orth_arrow_rounded(slide, waypoints, color, width_pt, style,
                            head if head != 'none' else None, corner_mm)


def orth_line(slide, waypoints, color, width_pt=0.8, style='solid', corner_mm=3.0):
    """Same as orth_arrow but with no arrowhead."""
    orth_arrow(slide, waypoints, color, width_pt=width_pt,
               style=style, head='none', corner_mm=corner_mm)


# ── Convenience: single-segment ──────────────────────────────────────────────

def arrow_right(slide, x, y, length, color, **kw):
    orth_arrow(slide, [(x, y), (x + int(length), y)], color, corner_mm=0, **kw)

def arrow_left(slide, x, y, length, color, **kw):
    orth_arrow(slide, [(x, y), (x - int(length), y)], color, corner_mm=0, **kw)

def arrow_down(slide, x, y, length, color, **kw):
    orth_arrow(slide, [(x, y), (x, y + int(length))], color, corner_mm=0, **kw)

def arrow_up(slide, x, y, length, color, **kw):
    orth_arrow(slide, [(x, y), (x, y - int(length))], color, corner_mm=0, **kw)


# ── Convenience: L-shapes (2 segments, rounded corner) ───────────────────────

def arrow_L_h(slide, x1, y1, corner_x, corner_y, color, corner_mm=3.0, **kw):
    """Horizontal first (y stays at y1), then vertical to corner_y."""
    orth_arrow(slide,
               [(x1, y1), (corner_x, y1), (corner_x, corner_y)],
               color, corner_mm=corner_mm, **kw)

def arrow_L_v(slide, x1, y1, corner_x, corner_y, color, corner_mm=3.0, **kw):
    """Vertical first (x stays at x1), then horizontal to corner_x."""
    orth_arrow(slide,
               [(x1, y1), (x1, corner_y), (corner_x, corner_y)],
               color, corner_mm=corner_mm, **kw)


# ─────────────────────────────────────────────────────────────────────────────
# Decorative markers
# ─────────────────────────────────────────────────────────────────────────────

def diamond_dot(slide, x, y, size_mm=1.8, color=None):
    """Small filled square used as line junction / terminator marker."""
    color = color or BT.PRIMARY_500_HEX
    sz = Mm(size_mm)
    s = slide.shapes.add_shape(1, int(x), int(y), int(sz), int(sz))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    return s


def label_flank_line(slide, cx, y, label, color,
                     line_half_mm=14, line_width_pt=0.8,
                     dot_size_mm=1.8, label_sz=9):
    """Renders:  ──── ◆ [label] ◆ ────  centered at (cx, y)."""
    half     = Mm(line_half_mm)
    dot      = Mm(dot_size_mm)
    dot_half = dot // 2
    label_h  = Mm(6)
    label_y  = y - label_h // 2
    label_w  = Mm(len(label) * 3.0)

    left_end   = int(cx - label_w / 2 - dot)
    right_start = int(cx + label_w / 2 + dot)

    c1 = _conn(slide, cx - half, y, left_end, y)
    _style_line(c1, color, line_width_pt)
    c2 = _conn(slide, right_start, y, cx + half, y)
    _style_line(c2, color, line_width_pt)

    diamond_dot(slide, left_end, y - dot_half, dot_size_mm, color)
    diamond_dot(slide, right_start - dot, y - dot_half, dot_size_mm, color)

    _txb(slide, label,
         l=int(cx - label_w / 2), t=label_y,
         w=int(label_w), h=label_h,
         sz=label_sz, bold=True, color=color,
         align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# Node components
# ─────────────────────────────────────────────────────────────────────────────

BOX_STYLES = {
    'default':   (BT.WHITE_HEX,         BT.NEUTRAL_200_HEX,   BT.NEUTRAL_400_HEX,   BT.NEUTRAL_900_HEX),
    'primary':   (BT.PRIMARY_100_HEX,   BT.PRIMARY_500_HEX,   BT.PRIMARY_500_HEX,   BT.NEUTRAL_900_HEX),
    'teal':      (BT.CARD_TEAL_BG,      BT.TEAL_HEX,          BT.TEAL_HEX,          BT.NEUTRAL_900_HEX),
    'success':   (BT.NEUTRAL_100_HEX,   BT.SUCCESS_HEX,       BT.SUCCESS_HEX,       BT.NEUTRAL_900_HEX),
    'secondary': (BT.SECONDARY_100_HEX, BT.SECONDARY_500_HEX, BT.SECONDARY_500_HEX, BT.NEUTRAL_900_HEX),
    'orange':    (BT.CARD_ORANGE_BG,    BT.WARNING_HEX,       BT.WARNING_HEX,       BT.NEUTRAL_900_HEX),
    'purple':    (BT.CARD_PURPLE_BG,    BT.PURPLE_HEX,        BT.PURPLE_HEX,        BT.NEUTRAL_900_HEX),
    'danger':    ('#FFF2F2',            BT.DANGER_HEX,        BT.DANGER_HEX,        BT.NEUTRAL_900_HEX),
    'dark':      (BT.NEUTRAL_900_HEX,   BT.NEUTRAL_900_HEX,   BT.SECONDARY_500_HEX, BT.WHITE_HEX),
}


def oval(slide, x, y, size_mm, color):
    """Circular badge (large-radius rounded rect)."""
    sz = Mm(size_mm)
    r  = size_mm / 2
    s  = slide.shapes.add_shape(5, int(x), int(y), int(sz), int(sz))
    _apply_radius(s, r, int(sz), int(sz))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    return s


def capability_box(slide, x, y, w, h, style='default',
                   badge_label='', num_label='', title='', body='',
                   title_sz=12, body_sz=9.5, badge_size_mm=8.5):
    """
    Capability node card:  [● badge]  num_label / title / body

    style: key from BOX_STYLES or 4-tuple (bg, border, accent, text_color)
    """
    if isinstance(style, str):
        bg, border, accent, txt = BOX_STYLES.get(style, BOX_STYLES['default'])
    else:
        bg, border, accent, txt = style

    _card(slide, l=x, t=y, w=w, h=h, bg=bg, border=border)

    PAD_L = Mm(4.5); PAD_T = Mm(4.5)
    BADGE = Mm(badge_size_mm); GAP = Mm(3)

    bx = x + PAD_L; by = y + PAD_T
    oval(slide, bx, by, badge_size_mm, accent)
    _txb(slide, badge_label, l=bx, t=by, w=BADGE, h=BADGE,
         sz=6.5, bold=True, color=BT.WHITE_HEX,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=False)

    label_x = bx + BADGE + GAP; label_w = w - PAD_L - BADGE - GAP - Mm(3)
    _txb(slide, num_label, l=label_x, t=by + Mm(0.5), w=label_w, h=Mm(7),
         sz=7, bold=True, color=accent)

    title_y = by + BADGE + Mm(2); content_w = w - PAD_L * 2
    _txb(slide, title, l=x + PAD_L, t=title_y, w=content_w, h=Mm(8),
         sz=title_sz, bold=True, color=txt)

    body_y = title_y + Mm(8); body_h = h - (body_y - y) - Mm(2.5)
    body_color = BT.NEUTRAL_200_HEX if bg == BT.NEUTRAL_900_HEX else BT.NEUTRAL_700_HEX
    _txb(slide, body, l=x + PAD_L, t=body_y, w=content_w, h=body_h,
         sz=body_sz, color=body_color)


def group_frame(slide, x, y, w, h, label='', accent=None,
                bg=None, dashed=True, radius_mm=5, label_sz=9):
    """Group container with optional dashed border and section label."""
    accent = accent or BT.PRIMARY_500_HEX
    bg     = bg     or BT.PRIMARY_100_HEX

    s = slide.shapes.add_shape(5, int(x), int(y), int(w), int(h))
    _apply_radius(s, radius_mm, int(w), int(h))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(bg)
    s.line.color.rgb = _rgb(accent); s.line.width = Pt(1.2)
    if dashed:
        spPr = s._element.find(qn('p:spPr'))
        ln   = spPr.find(qn('a:ln'))
        if ln is not None:
            etree.SubElement(ln, qn('a:prstDash')).set('val', 'dash')
    if label:
        _txb(slide, label, l=x, t=y + Mm(3), w=w, h=Mm(7),
             sz=label_sz, bold=True, color=accent, align=PP_ALIGN.CENTER)
    return s


def _add_drop_shadow(shape, blur_pt=6.0, dist_pt=2.0, dir_deg=90, alpha_pct=10):
    """Add a subtle outer drop shadow (floating card effect)."""
    sp   = shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    existing = spPr.find(qn('a:effectLst'))
    if existing is not None:
        spPr.remove(existing)
    eff  = etree.SubElement(spPr, qn('a:effectLst'))
    shdw = etree.SubElement(eff, qn('a:outerShdw'))
    shdw.set('blurRad',       str(int(Pt(blur_pt))))
    shdw.set('dist',          str(int(Pt(dist_pt))))
    shdw.set('dir',           str(int(dir_deg * 60000)))
    shdw.set('rotWithShape',  '0')
    clr = etree.SubElement(shdw, qn('a:srgbClr'))
    clr.set('val', '000000')
    etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha_pct * 1000)))


def node_card(slide, x, y, w, h, state='default', radius_mm=4):
    """
    IO boundary node (start/end point in a flow diagram).

    state='default'  : white bg + light gray border + drop shadow
                       → accessible start/end point
    state='inactive' : NEUTRAL_100 gray bg + gray border, no shadow
                       → disabled / blocked / not yet developed
    """
    if state == 'inactive':
        bg, shadow = BT.NEUTRAL_100_HEX, False
    else:
        bg, shadow = BT.WHITE_HEX, True
    s = slide.shapes.add_shape(5, int(x), int(y), int(w), int(h))
    _apply_radius(s, radius_mm, int(w), int(h))
    s.fill.solid(); s.fill.fore_color.rgb = _rgb(bg)
    s.line.color.rgb = _rgb(BT.NEUTRAL_200_HEX); s.line.width = Pt(0.75)
    if shadow:
        _add_drop_shadow(s)
    return s


def _set_shape_alpha_gradient(shape, color_hex, alpha_start_pct, alpha_end_pct=0,
                               angle_deg=90):
    """
    Replace a shape's fill with a 2-stop gradient that fades between two alpha values.
    alpha_*_pct: 0–100 (100 = fully opaque). angle_deg: 90 = top-to-bottom.
    """
    sp   = shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None:
        return

    for fill_tag in (qn('a:solidFill'), qn('a:gradFill'), qn('a:noFill'),
                     qn('a:pattFill'), qn('a:blipFill'), qn('a:grpFill')):
        el = spPr.find(fill_tag)
        if el is not None:
            spPr.remove(el)

    color_rgb = color_hex.upper().lstrip('#')
    a_start   = int(alpha_start_pct * 1000)
    a_end     = int(alpha_end_pct   * 1000)
    ang_ooxml = str(int(angle_deg * 60000))

    gf    = etree.Element(qn('a:gradFill'))
    gsLst = etree.SubElement(gf, qn('a:gsLst'))
    for pos_str, alpha_val in (('0', a_start), ('100000', a_end)):
        gs  = etree.SubElement(gsLst, qn('a:gs'))
        gs.set('pos', pos_str)
        clr = etree.SubElement(gs, qn('a:srgbClr'))
        clr.set('val', color_rgb)
        etree.SubElement(clr, qn('a:alpha')).set('val', str(alpha_val))
    lin = etree.SubElement(gf, qn('a:lin'))
    lin.set('ang', ang_ooxml); lin.set('scaled', '0')

    ln_el = spPr.find(qn('a:ln'))
    if ln_el is not None:
        spPr.insert(list(spPr).index(ln_el), gf)
    else:
        spPr.append(gf)


def _set_line_alpha(shape, alpha_pct):
    """Add alpha transparency to a shape's outline solidFill color."""
    sp   = shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    ln = spPr.find(qn('a:ln'))
    if ln is None: return
    sf = ln.find(qn('a:solidFill'))
    if sf is None: return
    clr = sf.find(qn('a:srgbClr'))
    if clr is None: return
    ex = clr.find(qn('a:alpha'))
    if ex is not None:
        clr.remove(ex)
    etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha_pct * 1000)))


def layer_frame(slide, x, y, w, h, label='', accent=None,
                bg_color=None, fill_alpha=32, angle_deg=90,
                dashed=True, border_alpha=50, border_pt=0.7,
                radius_mm=6, label_sz=9):
    """
    Transparent layer container for 'tier' / 'group' semantics.

    Renders as a gradient that fades from `fill_alpha`% at the dense end to
    fully transparent, plus a semi-transparent dashed border.  Visually
    lighter than solid BOX_STYLES cards, making the layer/element hierarchy
    immediately readable.

    fill_alpha   : opacity at gradient start (0–100 %)
    angle_deg    : 90 = top-to-bottom, 0 = left-to-right
    border_alpha : border line opacity (0–100 %)
    """
    accent   = accent   or BT.PRIMARY_500_HEX
    bg_color = bg_color or BT.PRIMARY_100_HEX

    s = slide.shapes.add_shape(5, int(x), int(y), int(w), int(h))
    _apply_radius(s, radius_mm, int(w), int(h))

    # Border (python-pptx creates the ln element here)
    s.line.color.rgb = _rgb(accent)
    s.line.width     = Pt(border_pt)
    if dashed:
        spPr = s._element.find(qn('p:spPr'))
        ln   = spPr.find(qn('a:ln'))
        if ln is not None:
            etree.SubElement(ln, qn('a:prstDash')).set('val', 'lgDash')
    _set_line_alpha(s, border_alpha)

    # Gradient fill fading to transparent
    _set_shape_alpha_gradient(s, bg_color, fill_alpha, 0, angle_deg)

    if label:
        _txb(slide, label, l=x, t=y + Mm(3), w=w, h=Mm(7),
             sz=label_sz, bold=True, color=accent, align=PP_ALIGN.CENTER)
    return s


def data_bar(slide, x, y, w, h,
             badge_label='', badge_color=None,
             title='', items_str='', title_sz=10, items_sz=9):
    """Full-width horizontal bar: [● badge] Title    item1 | item2 | ..."""
    badge_color = badge_color or BT.WARNING_HEX
    _card(slide, l=x, t=y, w=w, h=h, bg=BT.CARD_ORANGE_BG)
    badge_mm = 9
    bx = x + Mm(3); by = y + (h - Mm(badge_mm)) // 2
    oval(slide, bx, by, badge_mm, badge_color)
    _txb(slide, badge_label, l=bx, t=by, w=Mm(badge_mm), h=Mm(badge_mm),
         sz=6, bold=True, color=BT.WHITE_HEX,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=False)
    tx = x + Mm(14)
    _txb(slide, title, l=tx, t=y + Mm(2), w=Mm(80), h=Mm(7),
         sz=title_sz, bold=True, color=badge_color)
    _txb(slide, items_str, l=tx, t=y + Mm(8), w=w - (tx - x) - Mm(3), h=Mm(7),
         sz=items_sz, color=BT.NEUTRAL_700_HEX)


def summary_footer(slide, x, y, w, h, sections):
    """
    White footer bar divided into sections.
    sections: [{'dot_color': hex, 'label': str, 'body': str}, ...]
    """
    _card(slide, l=x, t=y, w=w, h=h, bg=BT.WHITE_HEX, border=BT.NEUTRAL_200_HEX)
    n = len(sections); sec_w = w // n
    for i, sec in enumerate(sections):
        sx = x + i * sec_w
        if i > 0:
            _rect(slide, l=sx, t=y + Mm(1.5), w=Mm(0.3), h=h - Mm(3),
                  fill=BT.NEUTRAL_200_HEX)
        dot_color = sec.get('dot_color', BT.PRIMARY_500_HEX)
        _rect(slide, l=sx + Mm(4), t=y + h // 2 - Mm(3),
              w=Mm(3.5), h=Mm(3.5), fill=dot_color, radius_mm=1.5)
        _txb(slide, sec.get('label', ''),
             l=sx + Mm(9), t=y + Mm(2), w=sec_w - Mm(12), h=Mm(5),
             sz=8, bold=True, color=dot_color)
        _txb(slide, sec.get('body', ''),
             l=sx + Mm(4), t=y + Mm(6.5), w=sec_w - Mm(7), h=h - Mm(7),
             sz=8.5, color=BT.NEUTRAL_700_HEX)
